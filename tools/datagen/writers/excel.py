"""Excel writers: supplier master, BOM master, inventory snapshot, pricelist.

Uses openpyxl. Writes merged headers and formulas on the BOM sheet so the
ingest pipeline exercises its 'bad table parsing' defences.
"""
from __future__ import annotations

import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from tools.datagen.entities import Pop, Supplier
from tools.datagen.mess import mixed_units_price, pick_name_variant
from tools.datagen.realism_log import RealismLog

_HEADER_FILL = PatternFill("solid", fgColor="DDE6F0")
_HEADER_FONT = Font(bold=True)
_CENTER = Alignment(horizontal="center", vertical="center")


def _apply_header(cell) -> None:
    cell.fill = _HEADER_FILL
    cell.font = _HEADER_FONT
    cell.alignment = _CENTER


def write_supplier_master(path: Path, pop: Pop, extras: list[Supplier],
                          rng: random.Random, log: RealismLog) -> Path:
    """Supplier master: multi-sheet workbook (Tier-1, Tier-2, Addresses).

    Deliberately mixes in some ``extras`` (near-duplicates with same tax_id)
    appended at random positions so an ingest run can't trivially spot them.
    """
    wb = Workbook()
    ws1 = wb.active
    assert ws1 is not None
    ws1.title = "Tier-1 Suppliers"
    ws2 = wb.create_sheet("Tier-2 Suppliers")
    ws3 = wb.create_sheet("Addresses")

    cols = ["Supplier ID", "Name", "Country", "City", "Tax ID",
            "Established", "Employees",
            "Annual Capacity (units)", "Capabilities", "Certifications",
            "Primary Currency", "Last Verified"]
    for i, c in enumerate(cols, start=1):
        _apply_header(ws1.cell(row=1, column=i, value=c))
        _apply_header(ws2.cell(row=1, column=i, value=c))

    def row_for(s: Supplier) -> list:
        missing = set(s.missing_fields)
        last_verified = "2022-01-15" if s.stale else "2026-07-10"
        return [
            s.id,
            pick_name_variant(rng, s, log, context=f"supplier_master:{s.id}"),
            s.country_name,
            s.city if "city" not in missing else None,
            s.tax_id if "tax_id" not in missing else None,
            s.established_year,
            s.employees,
            s.annual_capacity_units,
            "; ".join(s.capabilities),
            "; ".join(s.certifications),
            s.primary_currency if "primary_currency" not in missing else None,
            last_verified,
        ]

    # Interleave the near-dup extras into the main lists at random indices.
    tier1 = [s for s in pop.suppliers if s.tier == 1]
    tier2 = [s for s in pop.suppliers if s.tier == 2]
    for e in extras:
        bucket = tier1 if e.tier == 1 else tier2
        bucket.insert(rng.randint(0, len(bucket)), e)

    for r, s in enumerate(tier1, start=2):
        for c, v in enumerate(row_for(s), start=1):
            ws1.cell(row=r, column=c, value=v)
    for r, s in enumerate(tier2, start=2):
        for c, v in enumerate(row_for(s), start=1):
            ws2.cell(row=r, column=c, value=v)

    # Addresses sheet: one row per supplier, with the full postal string.
    for i, c in enumerate(["Supplier ID", "Legal Name", "Address",
                           "Country Code"], start=1):
        _apply_header(ws3.cell(row=1, column=i, value=c))
    for r, s in enumerate(pop.suppliers + extras, start=2):
        if "address" in s.missing_fields:
            addr = None
            log.event("supplier.missing_field_detail",
                      f"{s.id}: address blank in supplier_master.xlsx")
        else:
            addr = s.address
        ws3.cell(row=r, column=1, value=s.id)
        ws3.cell(row=r, column=2, value=s.canonical_name)
        ws3.cell(row=r, column=3, value=addr)
        ws3.cell(row=r, column=4, value=s.country_code)

    for ws in (ws1, ws2, ws3):
        for i in range(1, len(cols) + 1):
            ws.column_dimensions[get_column_letter(i)].width = 20

    wb.save(path)
    return path


def write_bom_master(path: Path, pop: Pop, rng: random.Random,
                     log: RealismLog) -> Path:
    """Each sheet = one tier-1 supplier's BOMs. The header for each SKU
    block uses a MERGED cell above three sub-columns (component / qty /
    unit cost) and a formula column for extended cost. This is the kind of
    header shape that trips naive table extraction.

    One sheet is deliberately malformed: the merged header is misaligned
    with its sub-columns. Ingest should fail extraction on it.
    """
    wb = Workbook()
    wb.remove(wb.active)  # drop default

    tier1 = [s for s in pop.suppliers if s.tier == 1]
    # Pick a sample so the workbook stays manageable.
    sample_suppliers = tier1[:12]
    comp_by_id = {c.id: c for c in pop.components}

    for s in sample_suppliers:
        safe = (s.id)[:28]  # Excel sheet-name limit
        ws = wb.create_sheet(title=safe)
        # Document header
        ws["A1"] = f"Bill of Materials — {s.canonical_name}"
        ws["A1"].font = Font(bold=True, size=14)
        ws.merge_cells("A1:F1")
        ws["A2"] = f"Supplier ID: {s.id}    |    Prepared: 2026-09-30"

        row = 4
        skus_for_s = [k for k in pop.skus if k.tier1_supplier_id == s.id][:6]
        for sku in skus_for_s:
            # Merged block header spanning Component/Qty/Unit cost/Ext cost.
            header_cell = ws.cell(
                row=row, column=1,
                value=f"SKU: {sku.id}  ({sku.garment_type}, "
                      f"{sku.color_name}, size {sku.size})")
            header_cell.font = Font(bold=True)
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)

            # Sub-header
            sub = ["Component ID", "Description", "Qty / unit", "Unit cost (USD)",
                   "Extended cost"]
            for i, h in enumerate(sub, start=1):
                _apply_header(ws.cell(row=row + 1, column=i, value=h))

            for j, (cid, qty) in enumerate(sku.bom):
                r = row + 2 + j
                comp = comp_by_id.get(cid)
                ws.cell(row=r, column=1, value=cid)
                ws.cell(row=r, column=2, value=comp.description if comp else "unknown")
                ws.cell(row=r, column=3, value=qty)
                ws.cell(row=r, column=4, value=comp.unit_cost_usd if comp else 0.0)
                # FORMULA: extended cost = qty * unit_cost
                ws.cell(row=r, column=5, value=f"=C{r}*D{r}")

            total_row = row + 2 + len(sku.bom)
            ws.cell(row=total_row, column=4, value="Total:").font = Font(bold=True)
            start_cell = row + 2
            end_cell = total_row - 1
            sum_cell = ws.cell(row=total_row, column=5,
                               value=f"=SUM(E{start_cell}:E{end_cell})")
            sum_cell.font = Font(bold=True)
            row = total_row + 2

        for i in range(1, 6):
            ws.column_dimensions[get_column_letter(i)].width = 24

    # One deliberately malformed BOM sheet.
    if sample_suppliers:
        ws_bad = wb.create_sheet(title="BOM-MALFORMED")
        ws_bad["A1"] = "BOM — formatting error (do not use)"
        ws_bad.merge_cells("A1:C1")
        # Merged header lies about the number of sub-columns:
        ws_bad["A3"] = "Component block"
        ws_bad.merge_cells("A3:D3")   # claims 4 cols
        # But only 3 sub-headers below:
        ws_bad["A4"] = "ID"
        ws_bad["B4"] = "Qty"
        ws_bad["C4"] = "Unit cost"
        ws_bad["A5"] = "CMP-FAB-0001"
        ws_bad["B5"] = 1.4
        ws_bad["C5"] = 3.20
        log.event("table.malformed_on_purpose",
                  "BOM-MALFORMED sheet: merged header spans 4 cols but only "
                  "3 sub-headers; extraction should fail or report low confidence.")

    wb.save(path)
    return path


def write_inventory_snapshot(path: Path, pop: Pop, rng: random.Random,
                              log: RealismLog) -> Path:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Inventory 2026-09-30"
    headers = ["SKU", "Plant ID", "Country", "On-hand units",
               "In-transit units", "Buffer days of demand"]
    for i, h in enumerate(headers, start=1):
        _apply_header(ws.cell(row=1, column=i, value=h))
    plant_by_id = {p.id: p for p in pop.plants}
    for r, sku in enumerate(pop.skus, start=2):
        plant = plant_by_id.get(sku.plant_id)
        onhand = rng.randint(0, 12_000)
        intransit = rng.randint(0, 4_000)
        buffer_days = rng.choice([3, 5, 7, 10, 14, 21, 28])
        ws.cell(row=r, column=1, value=sku.id)
        ws.cell(row=r, column=2, value=sku.plant_id)
        ws.cell(row=r, column=3, value=plant.country_code if plant else "")
        ws.cell(row=r, column=4, value=onhand)
        ws.cell(row=r, column=5, value=intransit)
        ws.cell(row=r, column=6, value=buffer_days)
    for i in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 20
    wb.save(path)
    return path


def write_pricelist(path: Path, pop: Pop, rng: random.Random,
                    log: RealismLog) -> Path:
    """Multi-sheet price list by season. Prices in USD. Some prices in the
    AW26 sheet are deliberately re-denominated in local currency and
    MIS-LABELLED as USD — the ingest's unit validator should catch this.
    """
    wb = Workbook()
    wb.remove(wb.active)
    supplier_by_id = {s.id: s for s in pop.suppliers}

    for season in ("SS26", "AW26", "SS27"):
        ws = wb.create_sheet(title=f"Prices {season}")
        headers = ["SKU", "Garment", "Fit", "Color", "Size",
                   "Retail (USD)", "Target cost (USD)",
                   "Tier-1 supplier"]
        for i, h in enumerate(headers, start=1):
            _apply_header(ws.cell(row=1, column=i, value=h))
        skus_season = [k for k in pop.skus if k.season == season]
        for r, sku in enumerate(skus_season, start=2):
            s = supplier_by_id.get(sku.tier1_supplier_id)
            ccy = s.primary_currency if s else "USD"
            # Mess injection concentrated on AW26.
            if season == "AW26":
                cost_val, _ = mixed_units_price(
                    rng, sku.cost_target_usd, ccy, log,
                    context=f"pricelist AW26:{sku.id}",
                )
            else:
                cost_val = sku.cost_target_usd
            ws.cell(row=r, column=1, value=sku.id)
            ws.cell(row=r, column=2, value=sku.garment_type)
            ws.cell(row=r, column=3, value=sku.fit)
            ws.cell(row=r, column=4, value=sku.color_name)
            ws.cell(row=r, column=5, value=sku.size)
            ws.cell(row=r, column=6, value=sku.retail_price_usd)
            ws.cell(row=r, column=7, value=cost_val)
            ws.cell(row=r, column=8, value=s.canonical_name if s else "")
        for i in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(i)].width = 18

    wb.save(path)
    return path


def write_supplier_financial_statement(path: Path, s: Supplier,
                                       rng: random.Random,
                                       log: RealismLog) -> Path:
    """Per-supplier simplified P&L + balance sheet. Numeric track food."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Financials FY2026"

    ws["A1"] = f"Financial Statement — {s.canonical_name}"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:D1")
    ws["A2"] = f"Supplier ID: {s.id}"
    ws["A3"] = f"Reporting currency: {s.primary_currency}"
    ws["A4"] = "FY: Apr-2025 to Mar-2026"

    # P&L
    ws["A6"] = "Profit & Loss"
    ws["A6"].font = Font(bold=True)
    headers = ["Line item", "Amount", "YoY", "Notes"]
    for i, h in enumerate(headers, start=1):
        _apply_header(ws.cell(row=7, column=i, value=h))
    revenue = int(s.annual_capacity_units * rng.uniform(0.4, 0.95)
                  * rng.uniform(4, 18))
    cogs = int(revenue * rng.uniform(0.62, 0.82))
    gm = revenue - cogs
    opex = int(revenue * rng.uniform(0.09, 0.18))
    ebit = gm - opex
    tax = max(0, int(ebit * 0.22))
    pnl = [
        ("Revenue", revenue, "+6.2%", ""),
        ("Cost of goods sold", cogs, "+5.1%", "incl. raw materials"),
        ("Gross margin", gm, "+9.3%", "=Revenue-COGS"),
        ("Operating expenses", opex, "+4.0%", ""),
        ("EBIT", ebit, "+14.8%", ""),
        ("Tax", tax, "", ""),
        ("Net income", ebit - tax, "+15.1%", ""),
    ]
    for r, (name, val, yoy, note) in enumerate(pnl, start=8):
        ws.cell(row=r, column=1, value=name)
        ws.cell(row=r, column=2, value=val)
        ws.cell(row=r, column=3, value=yoy)
        ws.cell(row=r, column=4, value=note)

    # Balance sheet summary
    ws["A17"] = "Balance Sheet (summary)"
    ws["A17"].font = Font(bold=True)
    for i, h in enumerate(["Line item", "Amount"], start=1):
        _apply_header(ws.cell(row=18, column=i, value=h))
    assets = int(revenue * rng.uniform(0.55, 1.2))
    liab = int(assets * rng.uniform(0.3, 0.65))
    bs = [
        ("Total assets", assets),
        ("Current assets", int(assets * 0.4)),
        ("Non-current assets", int(assets * 0.6)),
        ("Total liabilities", liab),
        ("Equity", assets - liab),
    ]
    for r, (name, val) in enumerate(bs, start=19):
        ws.cell(row=r, column=1, value=name)
        ws.cell(row=r, column=2, value=val)

    for i in range(1, 5):
        ws.column_dimensions[get_column_letter(i)].width = 22
    wb.save(path)
    return path
