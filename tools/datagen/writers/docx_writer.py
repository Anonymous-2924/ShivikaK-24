"""DOCX writers: policies, supplier correspondence, incident post-mortems."""
from __future__ import annotations

import random
from pathlib import Path

from docx import Document

from tools.datagen.entities import Supplier
from tools.datagen.mess import pick_name_variant
from tools.datagen.realism_log import RealismLog


def write_policy(path: Path, title: str, body_paragraphs: list[str]) -> Path:
    doc = Document()
    doc.add_heading(title, level=1)
    for p in body_paragraphs:
        doc.add_paragraph(p)
    doc.save(path)
    return path


def write_correspondence(path: Path, s: Supplier, rng: random.Random,
                          log: RealismLog) -> Path:
    doc = Document()
    name_in_doc = pick_name_variant(rng, s, log, context=f"email:{s.id}")
    doc.add_heading(f"Email thread with {name_in_doc}", level=1)

    turns = [
        ("Procurement <proc@grain-apparel.co>",
         "Subject: PO 10231 — fabric delay",
         "Hi, we have not yet received the shipping documents for PO 10231 "
         "covering the AW26 outer-shell fabric. Could you confirm whether "
         "the dye lot has been released?"),
        (f"Sales <orders@{s.canonical_name.lower().replace(' ', '-')}.com>",
         "Re: PO 10231 — fabric delay",
         rng.choice([
             "Apologies — the dye house has a 4-day backlog due to a boiler "
             "outage on 12 Sep. We'll courier the test swatch tomorrow.",
             "Please note the fabric mill (our tier-2 vendor) advised a "
             "4-week lead-time extension due to monsoon flooding affecting "
             "yarn delivery.",
             "We will prioritise this PO once we complete the Buyer-A order "
             "running on Line 3. Expected release by end of next week.",
         ])),
    ]
    for who, subj, body in turns:
        doc.add_paragraph(f"From: {who}").italic = True
        doc.add_paragraph(subj).bold = True
        doc.add_paragraph(body)
        doc.add_paragraph("")

    doc.save(path)
    return path


def write_postmortem(path: Path, title: str, affected_supplier: Supplier,
                     rng: random.Random, log: RealismLog) -> Path:
    doc = Document()
    doc.add_heading(title, level=1)
    doc.add_paragraph(f"Affected supplier: {affected_supplier.canonical_name} "
                      f"({affected_supplier.id})")
    doc.add_paragraph(f"Country: {affected_supplier.country_name}")

    doc.add_heading("Timeline", level=2)
    incident = rng.choice([
        "boiler failure", "power outage",
        "small fire in warehouse zone", "chemical spill",
    ])
    for line in [
        f"Day 0: Incident reported at the dyeing unit — {incident}.",
        "Day 1: Production halt confirmed for approximately 5 lines.",
        "Day 3: Alternative sourcing requested for the affected SKUs.",
        "Day 7: Partial recovery; 60% capacity restored.",
        "Day 14: Full recovery for Tier-1 operations; "
        "tier-2 fabric mill still catching up.",
    ]:
        doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("Root cause", level=2)
    doc.add_paragraph(rng.choice([
        "Deferred maintenance on the primary boiler combined with a monsoon-season "
        "humidity spike that exceeded design tolerances.",
        "Overloaded single-source dyeing capacity with insufficient alternate-vendor "
        "qualification for navy and black lots.",
        "Fire-detection system in the chemical storage area was decommissioned "
        "during a scheduled upgrade and not re-commissioned on schedule.",
    ]))

    doc.save(path)
    return path
